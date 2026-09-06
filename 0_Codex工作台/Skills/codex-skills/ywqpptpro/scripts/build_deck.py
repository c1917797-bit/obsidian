"""从已确认的 JSON spec 构建趋势驱动的技术洞察演示文稿。"""
from __future__ import annotations

import argparse
import copy
import json
import os
import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

from content_rules import compose_background_text
from validate_spec import validate_spec_data
from validate_template_fidelity import validate_template_fidelity

try:
    from PIL import Image
except ImportError:  # pragma: no cover - surfaced at runtime with a clear error.
    Image = None

COVER_IDX   = 0
SUMMARY_IDX = 1
TECHLIST_IDX = 2
INSIGHT_IDX  = 3
THANKS_IDX   = 4

INSIGHT_SHAPE_IDS = {
    "title":             5,
    "reference":        13,
    "background_table": 10,
    "tech_detail_figure":  15,
    "tech_detail_figure2": 21,
    "tech_detail_text":    23,
    "experiment_figure":   17,
    "experiment_method":   25,
    "experiment_result":   26,
    "insight_table":       11,
    "left_card":            4,
    "right_card":           6,
}

LINE_SPACING_PCT  = 120000   # 1.2×
SINGLE_SPACING_PCT = 100000  # 1.0×
_a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
_r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_TEMP_IMAGE_PATHS = []

FONT_FAMILY = "Microsoft YaHei"
COLORS = {
    "red": "CC0000",
    "blue": "1B3A6B",
    "cyan": "00C9C8",
    "orange": "FF6B35",
    "text": "333333",
    "muted": "666666",
    "light": "F0F0F0",
    "pale_blue": "F4F7FB",
    "pale_orange": "FFF3ED",
    "white": "FFFFFF",
    "line": "D8DEE8",
    "gray": "A6A6A6",
}


# ── shape lookup ────────────────────────────────────────────────────────────

def find_shape_by_id(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError(f"shape_id {shape_id} not found")

def find_shape_by_id_safe(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    return None

def remove_shape(slide, shape_id):
    sh = find_shape_by_id_safe(slide, shape_id)
    if sh is not None:
        sh._element.getparent().remove(sh._element)


# ── line-spacing & autofit ──────────────────────────────────────────────────

def _set_line_spacing(text_frame, pct):
    for para in text_frame.paragraphs:
        pPr = para._p.find(qn("a:pPr"))
        if pPr is None:
            pPr = etree.SubElement(para._p, qn("a:pPr"))
            para._p.insert(0, pPr)
        lnSpc = pPr.find(qn("a:lnSpc"))
        if lnSpc is None:
            lnSpc = etree.SubElement(pPr, qn("a:lnSpc"))
        for ch in list(lnSpc):
            lnSpc.remove(ch)
        sp = etree.SubElement(lnSpc, qn("a:spcPct"))
        sp.set("val", str(pct))

def apply_line_spacing(text_frame, pct=LINE_SPACING_PCT):
    _set_line_spacing(text_frame, pct)

def apply_single_spacing(text_frame):
    _set_line_spacing(text_frame, SINGLE_SPACING_PCT)

def apply_line_spacing_cell(cell, pct=LINE_SPACING_PCT):
    _set_line_spacing(cell.text_frame, pct)

def enable_text_autofit(text_frame):
    """Set normAutofit so PowerPoint shrinks the font to fit rather than clip."""
    txBody = text_frame._txBody
    bodyPr = txBody.find(qn("a:bodyPr"))
    if bodyPr is None:
        return
    for tag in [qn("a:noAutofit"), qn("a:normAutofit"), qn("a:spAutoFit")]:
        for el in bodyPr.findall(tag):
            bodyPr.remove(el)
    etree.SubElement(bodyPr, qn("a:normAutofit"))


def disable_text_autofit(text_frame):
    """保留模板字号；容量不合格时应修改文案，而不是让 PowerPoint 静默缩字。"""
    tx_body = text_frame._txBody
    body_pr = tx_body.find(qn("a:bodyPr"))
    if body_pr is None:
        return
    for tag in (qn("a:noAutofit"), qn("a:normAutofit"), qn("a:spAutoFit")):
        for element in body_pr.findall(tag):
            body_pr.remove(element)
    etree.SubElement(body_pr, qn("a:noAutofit"))


def _clean_text_lines(value):
    lines = [line.strip() for line in str(value or "").splitlines()]
    return [line for line in lines if line]


def _compact_slide_text(value, max_chars):
    """为总览页压缩长事实描述，完整内容仍保留在预览和来源链接中。"""
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if len(text) <= max_chars:
        return text
    window = text[:max_chars]
    split_at = max(
        window.rfind("。"),
        window.rfind("；"),
        window.rfind("，"),
    )
    if split_at >= int(max_chars * 0.62):
        window = window[: split_at + 1]
    return window.rstrip("，；。 ") + "…"


def _background_text(tech):
    """把结构化背景组合为“现状—痛点—课题问题”的自然段。"""
    return compose_background_text(tech.get("background"))


# ── unified visual language ──────────────────────────────────────────────────

def _rgb(name_or_hex):
    value = COLORS.get(name_or_hex, name_or_hex).lstrip("#")
    return RGBColor.from_string(value.upper())


def _set_run_style(run, *, size=14, color="text", bold=False):
    run.font.name = FONT_FAMILY
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    r_pr = run._r.get_or_add_rPr()
    east_asia = r_pr.find(qn("a:ea"))
    if east_asia is None:
        east_asia = etree.SubElement(r_pr, qn("a:ea"))
    east_asia.set("typeface", FONT_FAMILY)


def _format_text_frame(
    text_frame,
    text,
    *,
    size=14,
    color="text",
    bold=False,
    align=PP_ALIGN.LEFT,
    margin=0.08,
    valign=MSO_ANCHOR.MIDDLE,
    line_spacing_pct=LINE_SPACING_PCT,
):
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(margin)
    text_frame.margin_right = Inches(margin)
    text_frame.margin_top = Inches(margin)
    text_frame.margin_bottom = Inches(margin)
    text_frame.vertical_anchor = valign
    lines = _clean_text_lines(text) or [""]
    for index, line in enumerate(lines):
        para = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        para.alignment = align
        para.text = line
        for run in para.runs:
            _set_run_style(run, size=size, color=color, bold=bold)
    _set_line_spacing(text_frame, line_spacing_pct)
    enable_text_autofit(text_frame)
    return text_frame


def _set_shape_fill(shape, fill_color=None, line_color=None, line_width=1):
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill_color)
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = _rgb(line_color)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()


def _add_textbox(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    size=14,
    color="text",
    bold=False,
    align=PP_ALIGN.LEFT,
    fill=None,
    line=None,
    rounded=False,
    margin=0.08,
    valign=MSO_ANCHOR.MIDDLE,
    line_spacing_pct=LINE_SPACING_PCT,
):
    if rounded or fill or line:
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(
            shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        _set_shape_fill(shape, fill, line)
    else:
        shape = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
    _format_text_frame(
        shape.text_frame,
        text,
        size=size,
        color=color,
        bold=bold,
        align=align,
        margin=margin,
        valign=valign,
        line_spacing_pct=line_spacing_pct,
    )
    return shape


def _add_bullets(
    slide,
    x,
    y,
    w,
    h,
    items,
    *,
    size=13,
    color="text",
    fill=None,
    line=None,
    margin=0.12,
):
    text = "\n".join(f"• {_text}" for _text in _clean_text_lines("\n".join(str(i) for i in items)))
    return _add_textbox(
        slide,
        x,
        y,
        w,
        h,
        text,
        size=size,
        color=color,
        fill=fill,
        line=line,
        margin=margin,
        valign=MSO_ANCHOR.TOP,
    )


def _add_page_title(slide, title, eyebrow):
    _add_textbox(
        slide,
        0.34,
        0.16,
        11.25,
        0.50,
        title,
        size=28,
        color="red",
        bold=True,
        margin=0,
        line_spacing_pct=SINGLE_SPACING_PCT,
    )
    _add_textbox(
        slide,
        11.15,
        0.31,
        1.65,
        0.22,
        eyebrow,
        size=10,
        color="muted",
        bold=True,
        align=PP_ALIGN.RIGHT,
        margin=0,
        line_spacing_pct=SINGLE_SPACING_PCT,
    )


def _add_insight_bar(slide, label, text, *, y=6.46):
    """复刻单项洞察页底部灰底红标观点栏。"""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.53),
        Inches(y),
        Inches(12.45),
        Inches(0.64),
    )
    _set_shape_fill(bar, "D9D9D9", None)
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1.34),
        Inches(y),
        Inches(0.018),
        Inches(0.64),
    )
    _set_shape_fill(divider, "B7B7B7", None)
    _add_textbox(
        slide,
        0.64,
        y + 0.04,
        0.60,
        0.56,
        label,
        size=14,
        color="red",
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
        line_spacing_pct=SINGLE_SPACING_PCT,
    )
    _add_textbox(
        slide,
        1.47,
        y + 0.04,
        11.28,
        0.56,
        text,
        size=14,
        color="muted",
        bold=False,
        margin=0,
        line_spacing_pct=LINE_SPACING_PCT,
    )


def _add_cropped_picture(slide, image_path, x, y, w, h):
    if not os.path.exists(image_path):
        raise FileNotFoundError(f"Image not found: {image_path}")
    prepared_path = _prepare_image_for_box(
        image_path, Inches(w), Inches(h)
    )
    return slide.shapes.add_picture(
        prepared_path,
        Inches(x),
        Inches(y),
        width=Inches(w),
        height=Inches(h),
    )


def _new_content_slide(prs):
    """从未修改的单项洞察模板页复制内容页，禁止从空白页重画。"""
    return duplicate_slide(prs, INSIGHT_IDX)


def _prepare_template_content_slide(
    slide,
    *,
    title,
    reference,
    background_label,
    background_text,
    left_heading,
    right_heading,
    insight_label,
    insight_text,
):
    """保留模板标题、双栏框、背景带、观点带和母版页脚，只清空正文占位内容。"""
    for shape_id in (
        INSIGHT_SHAPE_IDS["tech_detail_figure"],
        INSIGHT_SHAPE_IDS["tech_detail_figure2"],
        INSIGHT_SHAPE_IDS["tech_detail_text"],
        INSIGHT_SHAPE_IDS["experiment_figure"],
        INSIGHT_SHAPE_IDS["experiment_method"],
        INSIGHT_SHAPE_IDS["experiment_result"],
    ):
        remove_shape(slide, shape_id)

    title_shape = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["title"])
    replace_text_preserving_format(
        title_shape.text_frame,
        title,
        line_spacing_pct=SINGLE_SPACING_PCT,
        autofit=False,
    )
    auto_fit_title(title_shape, title)

    reference_shape = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["reference"])
    replace_text_preserving_format(
        reference_shape.text_frame,
        reference,
        line_spacing_pct=SINGLE_SPACING_PCT,
        autofit=False,
    )

    left_title = find_shape_by_id(slide, 7)
    right_title = find_shape_by_id(slide, 9)
    replace_text_preserving_format(
        left_title.text_frame,
        left_heading,
        line_spacing_pct=SINGLE_SPACING_PCT,
        autofit=False,
    )
    replace_text_preserving_format(
        right_title.text_frame,
        right_heading,
        line_spacing_pct=SINGLE_SPACING_PCT,
        autofit=False,
    )

    background_table = find_shape_by_id(
        slide, INSIGHT_SHAPE_IDS["background_table"]
    ).table
    replace_cell_text(background_table.rows[0].cells[0], background_label)
    replace_cell_text(background_table.rows[0].cells[1], background_text)

    insight_table = find_shape_by_id(
        slide, INSIGHT_SHAPE_IDS["insight_table"]
    ).table
    replace_cell_text(insight_table.rows[0].cells[0], insight_label)
    replace_cell_text(insight_table.rows[0].cells[1], insight_text)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb("white")
    return slide


def _add_fitted_picture(slide, image_path, x, y, w, h):
    """完整放入真实图片，按原比例缩放，不裁掉图例、坐标轴或架构边界。"""
    if not os.path.exists(image_path):
        raise FileNotFoundError(f"Image not found: {image_path}")
    left, top, width, height = _fit_box_to_image(
        image_path,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    return slide.shapes.add_picture(
        image_path,
        left,
        top,
        width=width,
        height=height,
    )


def _normalize_text_frame_font(text_frame):
    for para in text_frame.paragraphs:
        for run in para.runs:
            run.font.name = FONT_FAMILY
            r_pr = run._r.get_or_add_rPr()
            east_asia = r_pr.find(qn("a:ea"))
            if east_asia is None:
                east_asia = etree.SubElement(r_pr, qn("a:ea"))
            east_asia.set("typeface", FONT_FAMILY)


def normalize_deck_fonts(prs):
    """统一所有模板页和新增页的中英文字体，保留原有字号和颜色。"""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                _normalize_text_frame_font(shape.text_frame)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        _normalize_text_frame_font(cell.text_frame)


# ── text writing helpers ────────────────────────────────────────────────────

def _pick_style_run(text_frame):
    candidates = []
    for para in text_frame.paragraphs:
        for run in para.runs:
            candidates.append(run)
    if not candidates:
        return None
    def is_bold(r):
        try: return bool(r.font.bold)
        except: return False
    pool = [r for r in candidates if not is_bold(r)] or candidates
    pool.sort(key=lambda r: len(r.text or ""), reverse=True)
    return pool[0]._r


def replace_text_preserving_format(text_frame, new_text,
                                   line_spacing_pct=LINE_SPACING_PCT,
                                   autofit=True):
    if text_frame is None:
        return
    style_r   = _pick_style_run(text_frame)
    style_xml = copy.deepcopy(style_r) if style_r is not None else None
    body = text_frame._txBody
    for p in body.findall(f"{{{_a_ns}}}p"):
        body.remove(p)
    for line in (_clean_text_lines(new_text) or [""]):
        p_el = body.makeelement(qn("a:p"), {})
        body.append(p_el)
        if style_xml is not None:
            r_el = copy.deepcopy(style_xml)
            t_el = r_el.find(qn("a:t"))
            if t_el is not None:
                t_el.text = line
            p_el.append(r_el)
        else:
            r_el = p_el.makeelement(qn("a:r"), {})
            t_el = r_el.makeelement(qn("a:t"), {})
            t_el.text = line
            r_el.append(t_el)
            p_el.append(r_el)
    _set_line_spacing(text_frame, line_spacing_pct)
    if autofit:
        enable_text_autofit(text_frame)
    else:
        disable_text_autofit(text_frame)


def replace_cell_text(
    cell,
    new_text,
    line_spacing_pct=LINE_SPACING_PCT,
    *,
    autofit=True,
):
    replace_text_preserving_format(cell.text_frame, new_text,
                                   line_spacing_pct=line_spacing_pct,
                                   autofit=autofit)


def write_bold_prefix_paragraphs(text_frame, lines,
                                  line_spacing_pct=LINE_SPACING_PCT,
                                  *,
                                  autofit=True):
    """Write paragraphs where text before fullwidth '：' is bold.
    lines can be plain strings or ('BOLD'/'NORMAL', text) tuples.
    """
    style_r   = _pick_style_run(text_frame)
    style_xml = copy.deepcopy(style_r) if style_r is not None else None
    body = text_frame._txBody
    for p in body.findall(f"{{{_a_ns}}}p"):
        body.remove(p)

    def make_run(tmpl, content, bold):
        if tmpl is not None:
            r_el = copy.deepcopy(tmpl)
            rPr  = r_el.find(qn("a:rPr"))
            if rPr is None:
                rPr = etree.SubElement(r_el, qn("a:rPr"))
                r_el.insert(0, rPr)
            if bold:
                rPr.set("b", "1")
            else:
                rPr.attrib.pop("b", None)
            t_el = r_el.find(qn("a:t"))
            if t_el is None:
                t_el = etree.SubElement(r_el, qn("a:t"))
            t_el.text = content
        else:
            r_el = etree.Element(qn("a:r"))
            rPr  = etree.SubElement(r_el, qn("a:rPr"))
            if bold:
                rPr.set("b", "1")
            t_el = etree.SubElement(r_el, qn("a:t"))
            t_el.text = content
        return r_el

    cleaned_lines = []
    for line in lines:
        if isinstance(line, tuple):
            mode, content = line
            content = str(content or "").strip()
            if content:
                cleaned_lines.append((mode, content))
        else:
            content = str(line or "").strip()
            if content:
                cleaned_lines.append(content)

    for line in cleaned_lines:
        p_el = body.makeelement(qn("a:p"), {})
        body.append(p_el)
        if isinstance(line, tuple):
            mode, content = line
            p_el.append(make_run(style_xml, content, mode == "BOLD"))
        else:
            sep = "\uff1a" if "\uff1a" in line else (
                  ":" if ":" in line and line.index(":") < 28 else None)
            if sep:
                idx    = line.index(sep) + len(sep)
                prefix = line[:idx]
                rest   = line[idx:]
                if prefix: p_el.append(make_run(style_xml, prefix, True))
                if rest:   p_el.append(make_run(style_xml, rest,   False))
            else:
                p_el.append(make_run(style_xml, line, False))

    _set_line_spacing(text_frame, line_spacing_pct)
    if autofit:
        enable_text_autofit(text_frame)
    else:
        disable_text_autofit(text_frame)


# ── picture replacement (z-order safe, no shared-blob issue) ─────────────────

def _box_from_shape(shape):
    return shape.left, shape.top, shape.width, shape.height


def _union_box(shapes):
    left = min(sh.left for sh in shapes)
    top = min(sh.top for sh in shapes)
    right = max(sh.left + sh.width for sh in shapes)
    bottom = max(sh.top + sh.height for sh in shapes)
    return left, top, right - left, bottom - top


def _image_ratio(path):
    if Image is None:
        raise RuntimeError("Pillow is required for aspect-ratio-safe image placement.")
    with Image.open(path) as img:
        return img.width / img.height


def _fit_box_to_image(image_path, left, top, width, height):
    """在给定区域内按原图比例放大，不裁切图片内容。"""
    ratio = _image_ratio(image_path)
    box_ratio = float(width) / float(height)
    if ratio >= box_ratio:
        fitted_w = int(width)
        fitted_h = max(1, int(fitted_w / ratio))
        fitted_x = int(left)
        fitted_y = int(top + (height - fitted_h) / 2)
    else:
        fitted_h = int(height)
        fitted_w = max(1, int(fitted_h * ratio))
        fitted_x = int(left + (width - fitted_w) / 2)
        fitted_y = int(top)
    return fitted_x, fitted_y, fitted_w, fitted_h


def _prepare_image_for_box(image_path, width, height):
    """Center-crop image to the target box ratio so PowerPoint does not distort it."""
    if Image is None:
        raise RuntimeError("Pillow is required for aspect-ratio-safe image placement.")
    target_ratio = float(width) / float(height)
    img = Image.open(image_path).convert("RGB")
    src_ratio = img.width / img.height
    if abs(src_ratio - target_ratio) < 0.01:
        prepared = img
    elif src_ratio > target_ratio:
        new_w = int(img.height * target_ratio)
        left = max(0, (img.width - new_w) // 2)
        prepared = img.crop((left, 0, left + new_w, img.height))
    else:
        new_h = int(img.width / target_ratio)
        top = max(0, (img.height - new_h) // 2)
        prepared = img.crop((0, top, img.width, top + new_h))

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
    tmp.close()
    prepared.save(tmp.name)
    _TEMP_IMAGE_PATHS.append(tmp.name)
    return tmp.name


def _remove_shape_elements(shapes):
    for sh in shapes:
        if sh is not None and sh._element.getparent() is not None:
            sh._element.getparent().remove(sh._element)


def _add_picture_at(slide, anchor_shape, image_path, left, top, width, height,
                    remove_shapes):
    if not os.path.exists(image_path):
        raise FileNotFoundError(f"Image not found: {image_path}")
    spTree = anchor_shape._element.getparent()
    z_idx = list(spTree).index(anchor_shape._element)
    prepared_path = _prepare_image_for_box(image_path, width, height)
    _remove_shape_elements(remove_shapes)
    new_sh = slide.shapes.add_picture(
        prepared_path, left, top, width=width, height=height)
    new_el = new_sh._element
    spTree.remove(new_el)
    spTree.insert(z_idx, new_el)
    return new_sh


def replace_picture(slide, shape_id, new_image_path):
    """Replace picture while preserving z-order and image aspect ratio.

    Uses remove-then-add, but re-inserts the new shape element at the
    same spTree index as the original. The source image is center-cropped
    to the placeholder ratio before insertion, so PowerPoint never stretches it.
    Also creates a fresh image Part per call, avoiding the shared-blob
    issue that arises when slides are cloned and all reference the same Part.
    """
    pic = find_shape_by_id(slide, shape_id)
    left, top, width, height = _box_from_shape(pic)
    return _add_picture_at(slide, pic, new_image_path, left, top, width, height,
                           [pic])


def _coerce_image_list(value, *fallbacks):
    paths = []
    if isinstance(value, list):
        paths.extend(value)
    elif value:
        paths.append(value)
    paths.extend(f for f in fallbacks if f)
    seen = set()
    result = []
    for path in paths:
        if path and path not in seen:
            seen.add(path)
            result.append(path)
    return result


def _split_box_for_images(image_paths, left, top, width, height):
    if len(image_paths) <= 1:
        return [(left, top, width, height)]
    gap = int(min(width, height) * 0.035)
    ratios = [_image_ratio(path) for path in image_paths[:2]]
    side_ratio = ((width - gap) / 2) / height
    stack_ratio = width / ((height - gap) / 2)

    def penalty(target):
        return sum(abs(r - target) / max(r, target) for r in ratios)

    if penalty(stack_ratio) < penalty(side_ratio):
        half_h = int((height - gap) / 2)
        return [
            (left, top, width, half_h),
            (left, top + half_h + gap, width, height - half_h - gap),
        ]
    half_w = int((width - gap) / 2)
    return [
        (left, top, half_w, height),
        (left + half_w + gap, top, width - half_w - gap, height),
    ]


def replace_picture_group(
    slide,
    anchor_shape,
    image_paths,
    box,
    remove_shapes,
    *,
    preserve_content=False,
):
    image_paths = image_paths[:2]
    left, top, width, height = box
    boxes = _split_box_for_images(image_paths, left, top, width, height)
    spTree = anchor_shape._element.getparent()
    z_idx = list(spTree).index(anchor_shape._element)
    _remove_shape_elements(remove_shapes)
    inserted = []
    for path, (x, y, w, h) in zip(image_paths, boxes):
        if not os.path.exists(path):
            raise FileNotFoundError(f"Image not found: {path}")
        if preserve_content:
            x, y, w, h = _fit_box_to_image(path, x, y, w, h)
            prepared_path = path
        else:
            prepared_path = _prepare_image_for_box(path, w, h)
        new_sh = slide.shapes.add_picture(prepared_path, x, y, width=w, height=h)
        inserted.append(new_sh._element)
    for el in inserted:
        spTree.remove(el)
    for offset, el in enumerate(inserted):
        spTree.insert(z_idx + offset, el)


def _detail_layout_boxes(image_paths, image_box, text_box, *, min_text_fraction=0.30):
    """按图片比例分配图文区域，并为文字保留连续空间。

    返回 (figure_boxes, detail_text_box, layout_mode)；layout_mode 为
    "left_right"（图左文右）或 "top_bottom"（图上文下），供调用方决定
    文字框的垂直对齐方式。`min_text_fraction` 是文字区域相对于图文总高度
    的最小占比下限——图片实际渲染得比这个下限允许的还矮时，多出的空间
    会继续让给文字，而不是留成空白；只有图片本身需要更多高度时，才会
    被这个下限压缩。
    """
    left, top, width, _ = image_box
    content_bottom = text_box[1] + text_box[3]
    content_height = content_bottom - top
    gap = max(1, int(min(width, content_height) * 0.035))
    ratios = [_image_ratio(path) for path in image_paths[:2]]

    if len(ratios) == 1 and ratios[0] < 1.40:
        # 竖向图采用左图右文，避免图片上、下方形成大块空区。
        max_image_width = int(width * 0.54)
        image_height = int(content_height)
        image_width = int(image_height * ratios[0])
        if image_width > max_image_width:
            image_width = max_image_width
            image_height = int(image_width / ratios[0])
        image_top = int(top + (content_height - image_height) / 2)
        figure_boxes = [(left, image_top, image_width, image_height)]
        detail_text_box = (
            left + image_width + gap,
            top,
            width - image_width - gap,
            content_height,
        )
        return figure_boxes, detail_text_box, "left_right"

    min_text_height = int(content_height * min_text_fraction)
    max_image_height = max(1, content_height - gap - min_text_height)
    if len(ratios) == 1:
        image_height = min(max_image_height, int(width / ratios[0]))
        image_width = min(width, int(image_height * ratios[0]))
        figure_boxes = [
            (left + int((width - image_width) / 2), top, image_width, image_height)
        ]
    else:
        inner_width = width - gap
        image_height = min(max_image_height, int(inner_width / sum(ratios)))
        image_widths = [max(1, int(image_height * ratio)) for ratio in ratios]
        total_width = sum(image_widths) + gap
        if total_width > width:
            scale = (width - gap) / sum(image_widths)
            image_widths = [max(1, int(item * scale)) for item in image_widths]
            image_height = min(
                int(image_widths[index] / ratios[index])
                for index in range(len(ratios))
            )
        start_x = left + int((width - sum(image_widths) - gap) / 2)
        figure_boxes = [
            (start_x, top, image_widths[0], image_height),
            (
                start_x + image_widths[0] + gap,
                top,
                image_widths[1],
                image_height,
            ),
        ]

    detail_text_top = top + image_height + gap
    detail_text_box = (
        left,
        detail_text_top,
        width,
        content_bottom - detail_text_top,
    )
    return figure_boxes, detail_text_box, "top_bottom"


def reflow_figure_and_text(
    slide,
    anchor_shape,
    image_paths,
    image_box,
    text_shape,
    remove_shapes,
    *,
    text_box=None,
    min_text_fraction=0.30,
):
    """插入 1–2 张完整图片，并把文字框重新排布到图片旁边/下方。

    双图固定横向，宽度随原图比例变化；图片比文字框原有留出的空间矮时，
    多出的高度会让给文字，而不是留成空白。`text_box` 默认取自
    `text_shape` 当前的位置，调用方也可以传入一个扩大过下边界的
    合成 box（例如把模板卡片底部的留白也算进可用范围）。
    """
    image_paths = image_paths[:2]
    if text_box is None:
        text_box = _box_from_shape(text_shape)
    figure_boxes, detail_text_box, layout_mode = _detail_layout_boxes(
        image_paths, image_box, text_box, min_text_fraction=min_text_fraction
    )
    spTree = anchor_shape._element.getparent()
    z_idx = list(spTree).index(anchor_shape._element)
    _remove_shape_elements(remove_shapes)

    inserted = []
    for path, (x, y, w, h) in zip(image_paths, figure_boxes):
        if not os.path.exists(path):
            raise FileNotFoundError(f"Image not found: {path}")
        new_shape = slide.shapes.add_picture(path, x, y, width=w, height=h)
        inserted.append(new_shape._element)
    for element in inserted:
        spTree.remove(element)
    for offset, element in enumerate(inserted):
        spTree.insert(z_idx + offset, element)

    text_shape.left = int(detail_text_box[0])
    text_shape.top = int(detail_text_box[1])
    text_shape.width = max(1, int(detail_text_box[2]))
    text_shape.height = max(1, int(detail_text_box[3]))
    if layout_mode == "left_right":
        # 图片在左侧时，右侧文字整体垂直居中，避免顶对齐显得和图片不平衡。
        text_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    else:
        text_shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP


def cleanup_temp_images():
    for path in list(_TEMP_IMAGE_PATHS):
        try:
            os.remove(path)
        except OSError:
            pass
    _TEMP_IMAGE_PATHS.clear()


def validate_output_text_fit(deck_path):
    """在 Windows PowerPoint 中测量四个指定区域的真实文本边界。"""
    if os.name != "nt":
        print("[WARN] 非 Windows 环境，跳过 PowerPoint COM 文字边界复检。")
        return
    checker = Path(__file__).resolve().parent / "check_ppt_text_overflow.ps1"
    result = subprocess.run(
        [
            "powershell.exe",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(checker),
            "-Deck",
            str(deck_path),
        ],
        capture_output=True,
        text=True,
        errors="replace",
        check=False,
    )
    if result.stdout.strip():
        print(result.stdout.strip())
    if result.returncode == 2:
        print(
            "[WARN] PowerPoint COM 不可用；已保留生成前版面预算校验，"
            "仍需在视觉检查阶段逐页确认。"
        )
        return
    if result.returncode != 0:
        details = result.stderr.strip() or result.stdout.strip()
        raise ValueError(f"成品文字版面校验失败：\n{details}")


# ── title auto-shrink ────────────────────────────────────────────────────────

def auto_fit_title(title_shape, title_text):
    """Shrink font so the title stays on one line (threshold: 65 vis-width units)."""
    visible = sum(2 if ord(c) > 127 else 1 for c in title_text)
    if visible <= 65:
        return
    new_pt = max(13, int(22 * 65 / visible))
    for para in title_shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(new_pt)


def infer_reference_url(reference, explicit_url=None):
    if explicit_url:
        return explicit_url
    ref = reference or ""
    url_match = re.search(r"https?://[^\s,;]+", ref)
    if url_match:
        return url_match.group(0)
    arxiv_match = re.search(r"arxiv\s*:\s*(\d{4}\.\d{4,5})(?:v\d+)?", ref, re.I)
    if arxiv_match:
        return f"https://arxiv.org/abs/{arxiv_match.group(1)}"
    doi_match = re.search(r"doi\s*:\s*([^\s,;]+)", ref, re.I)
    if doi_match:
        return f"https://doi.org/{doi_match.group(1)}"
    return None


def apply_text_hyperlink(text_frame, url):
    if not url:
        return
    for para in text_frame.paragraphs:
        for run in para.runs:
            if run.text:
                run.hyperlink.address = url


# ── slide management ─────────────────────────────────────────────────────────

def duplicate_slide(prs, source_idx):
    source   = prs.slides[source_idx]
    new_sl   = prs.slides.add_slide(source.slide_layout)
    src_tree = source.shapes._spTree
    dst_tree = new_sl.shapes._spTree
    for shp in list(dst_tree):
        if shp.tag.split("}", 1)[-1] in ("sp","pic","graphicFrame","grpSp","cxnSp"):
            dst_tree.remove(shp)
    for shp in src_tree:
        if shp.tag.split("}", 1)[-1] in ("sp","pic","graphicFrame","grpSp","cxnSp"):
            dst_tree.append(copy.deepcopy(shp))
    src_part = source.part
    dst_part = new_sl.part
    for blip in dst_tree.iter(f"{{{_a_ns}}}blip"):
        embed = blip.get(f"{{{_r_ns}}}embed")
        if embed and embed in src_part.rels:
            rel     = src_part.rels[embed]
            new_rid = dst_part.relate_to(rel.target_part, rel.reltype)
            blip.set(f"{{{_r_ns}}}embed", new_rid)
    for hlink in dst_tree.iter(f"{{{_a_ns}}}hlinkClick"):
        rid = hlink.get(f"{{{_r_ns}}}id")
        if rid and rid in src_part.rels:
            rel = src_part.rels[rid]
            new_rid = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True) \
                      if rel.is_external else dst_part.relate_to(rel.target_part, rel.reltype)
            hlink.set(f"{{{_r_ns}}}id", new_rid)
    return new_sl


def reorder_slide_to(prs, slide, new_index):
    sldIdLst = prs.slides._sldIdLst
    target   = slide.part.partname
    for sldId in sldIdLst:
        rId = sldId.get(f"{{{_r_ns}}}id")
        if prs.part.related_part(rId).partname == target:
            sldIdLst.remove(sldId)
            sldIdLst.insert(new_index, sldId)
            return
    raise RuntimeError("Could not locate sldId for slide")


def remove_slide(prs, slide_idx):
    sldIdLst = prs.slides._sldIdLst
    sldId    = list(sldIdLst)[slide_idx]
    rId      = sldId.get(f"{{{_r_ns}}}id")
    prs.part.drop_rel(rId)
    sldIdLst.remove(sldId)


# ── per-slide population ─────────────────────────────────────────────────────

def populate_cover(slide, spec):
    title  = spec.get("title", "").replace("\n", " ")   # collapse to one line
    dept   = spec.get("department", "")
    author = spec.get("author", "")
    date   = spec.get("date", "")
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if "技术洞察" in txt and ("xxxx" in txt or txt.strip() == "技术洞察"):
            # Cover title: single spacing, bold
            replace_text_preserving_format(sh.text_frame, title,
                                           line_spacing_pct=SINGLE_SPACING_PCT,
                                           autofit=False)
            # Make every run bold
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    run.font.bold = True
                    run.font.size = Pt(32)
                    run.font.color.rgb = _rgb("red")
        elif "部门" in txt and "作者" in txt:
            replace_text_preserving_format(
                sh.text_frame,
                f"部门：{dept}\n作者：{author}\n日期：{date}",
                line_spacing_pct=SINGLE_SPACING_PCT, autofit=False)


def populate_summary(slide, spec):
    topic      = spec.get("topic", "")
    directions = spec.get("directions", [])
    overall    = spec.get("overall_summary", "")
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if "本次洞察主要围绕" in txt:
            direction_count = len(directions)
            replace_text_preserving_format(
                sh.text_frame,
                f"{topic}，关键在以下 {direction_count} 个方向：")
        elif "方向1：" in txt or "方向 1：" in txt:
            lines = []
            for direction in directions:
                link = (
                    direction.get("trend_link")
                    if isinstance(direction.get("trend_link"), dict)
                    else {}
                )
                stage_id = link.get("stage_id")
                stage_label = f"（{stage_id}）" if stage_id else ""
                lines.append(
                    f"{direction.get('direction_id', '')} "
                    f"{direction.get('name', '')}{stage_label}："
                    f"{direction.get('summary', '')}"
                )
            write_bold_prefix_paragraphs(sh.text_frame, lines)
        elif "总结：" in txt:
            replace_text_preserving_format(sh.text_frame, overall)


def populate_techlist(slide, spec):
    tbl_sh = next((sh for sh in slide.shapes if sh.has_table), None)
    if not tbl_sh:
        return
    plan = []
    for d in spec.get("directions", []):
        link = d.get("trend_link") if isinstance(d.get("trend_link"), dict) else {}
        stage_id = link.get("stage_id")
        category = d.get("category_label") or d.get("name", "")
        cat = f"[{stage_id}] {category}" if stage_id else category
        techs = d.get("technologies", [])
        for j, t in enumerate(techs):
            plan.append((cat if j == 0 else "", t.get("name",""), t.get("brief","")))
    tbl_el  = tbl_sh.table._tbl
    hr      = 1   # header rows
    cur     = len(tbl_sh.table.rows) - hr
    target  = len(plan) if plan else 6
    if target > cur:
        last = tbl_el.findall(qn("a:tr"))[-1]
        for _ in range(target - cur):
            tbl_el.append(copy.deepcopy(last))
    elif target < cur:
        for r in tbl_el.findall(qn("a:tr"))[hr + target:]:
            tbl_el.remove(r)
    table = tbl_sh.table
    # 模板首列默认跨行合并；趋势标签需要逐方向显示，先解除合并。
    for row_index in range(hr, len(table.rows)):
        cell = table.cell(row_index, 0)
        if cell.is_merge_origin:
            cell.split()
    data_row_height = max(0.44, min(0.72, 3.17 / max(1, target)))
    table.rows[0].height = Inches(0.55)
    for row_index in range(hr, len(table.rows)):
        table.rows[row_index].height = Inches(data_row_height)
    for i, (cat, name, brief) in enumerate(plan):
        row = table.rows[hr + i]
        replace_cell_text(row.cells[0], cat)
        replace_cell_text(row.cells[1], name)
        replace_cell_text(row.cells[2], brief)
        for column in range(3):
            cell = row.cells[column]
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    _set_run_style(
                        run,
                        size=13 if column != 2 else 12.5,
                        color="text",
                        bold=column in (0, 1),
                    )
    for i in range(len(plan), len(table.rows) - hr):
        for c in range(3):
            replace_cell_text(table.rows[hr + i].cells[c], "")


def decorate_summary_slide(slide, spec):
    """为模板摘要页补充标题和技术阶段到方向的逻辑图。"""
    _add_page_title(slide, "洞察摘要", "SUMMARY · LOGIC")
    directions = spec.get("directions", [])[:3]
    if not directions:
        return
    _add_textbox(
        slide,
        0.88,
        4.20,
        1.02,
        0.48,
        "技术脉络",
        size=12,
        color="blue",
        bold=True,
        margin=0,
    )
    gap = 0.30
    total_width = 10.35
    card_width = (total_width - gap * (len(directions) - 1)) / len(directions)
    for index, direction in enumerate(directions):
        x = 2.00 + index * (card_width + gap)
        link = (
            direction.get("trend_link")
            if isinstance(direction.get("trend_link"), dict)
            else {}
        )
        _add_textbox(
            slide,
            x,
            4.10,
            card_width,
            0.68,
            f"{link.get('stage_id', '')} → {direction.get('direction_id', '')}｜{direction.get('name', '')}",
            size=11.5,
            color="blue",
            bold=True,
            align=PP_ALIGN.CENTER,
            fill="pale_blue",
            line="gray",
            rounded=True,
            margin=0.06,
        )


def decorate_techlist_slide(slide):
    _add_page_title(slide, "技术清单", "MAP · TECHNOLOGY")


def add_trend_link_ribbon(slide, tech, direction=None):
    direction = direction or {}
    link = tech.get("trend_link") if isinstance(tech.get("trend_link"), dict) else {}
    direction_link = (
        direction.get("trend_link")
        if isinstance(direction.get("trend_link"), dict)
        else {}
    )
    stage_id = link.get("stage_id") or direction_link.get("stage_id") or "T?"
    claim = link.get("claim") or direction_link.get("claim") or ""
    text = f"{stage_id}｜{claim}"
    # 窄条放在页面右上角：右边与下方双栏框对齐（右边界 ≈12.92in），
    # 顶部卡在标题框底部（0.628in）与技术背景灰带顶部（0.944in）之间，
    # 几何上保证无论标题多长都不会被压住或压住标题。
    _add_textbox(
        slide,
        9.32,
        0.66,
        3.6,
        0.25,
        text,
        size=9,
        color="blue",
        bold=True,
        align=PP_ALIGN.RIGHT,
        fill="pale_blue",
        line="gray",
        margin=0.025,
        line_spacing_pct=SINGLE_SPACING_PCT,
    )


def populate_insight(slide, tech, direction=None):
    # ── Title (auto-shrink to one line, single spacing, bold preserved)
    title_text  = tech.get("title", "")
    title_shape = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["title"])
    replace_text_preserving_format(title_shape.text_frame, title_text,
                                   line_spacing_pct=SINGLE_SPACING_PCT,
                                   autofit=False)
    auto_fit_title(title_shape, title_text)

    # ── Reference (just below title, single spacing)
    ref_shape = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["reference"])
    ref       = tech.get("reference", "")
    replace_text_preserving_format(ref_shape.text_frame,
                                   f"参考： {ref}" if ref else "参考：",
                                   line_spacing_pct=SINGLE_SPACING_PCT,
                                   autofit=False)
    apply_text_hyperlink(
        ref_shape.text_frame,
        infer_reference_url(ref, tech.get("reference_url")),
    )
    add_trend_link_ribbon(slide, tech, direction)

    # ── Background table cell
    bg_tbl = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["background_table"])
    replace_cell_text(
        bg_tbl.table.rows[0].cells[1],
        _background_text(tech),
        autofit=False,
    )

    # ── Tech detail: 核心思想（人类可读整句）+ 关键技术（加粗名称的分点列表）
    det_sh = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["tech_detail_text"])
    det_lines = []
    core_idea = str(tech.get("tech_detail_summary") or "").strip()
    if core_idea:
        det_lines.append(f"核心思想：{core_idea}")
    det_lines.append(("BOLD", "关键技术："))
    for point in tech.get("tech_detail_points", []):
        point = str(point or "").strip()
        if point:
            det_lines.append(f"• {point}")
    write_bold_prefix_paragraphs(det_sh.text_frame, det_lines, autofit=False)

    # ── Tech-detail figures. One image adapts to its ratio; two stay horizontal.
    detail_figs = _coerce_image_list(
        tech.get("tech_detail_figures"),
        tech.get("tech_detail_figure"),
        tech.get("tech_detail_figure2"),
    )
    if not detail_figs:
        raise ValueError(f"{title_text}: missing tech detail figure")
    fig_slot_1 = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["tech_detail_figure"])
    fig_slot_2 = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["tech_detail_figure2"])
    reflow_figure_and_text(
        slide,
        fig_slot_1,
        detail_figs,
        _union_box([fig_slot_1, fig_slot_2]),
        det_sh,
        [fig_slot_1, fig_slot_2],
        min_text_fraction=0.40,
    )

    # ── Experiment method (bold header)
    meth_sh    = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["experiment_method"])
    meth_lines = []
    if tech.get("experiment_method_header"):
        meth_lines.append(("BOLD", tech["experiment_method_header"]))
    meth_lines.extend(tech.get("experiment_method_points", []))
    write_bold_prefix_paragraphs(
        meth_sh.text_frame,
        meth_lines,
        autofit=False,
    )

    # ── Experiment results: 实验总结（人类可读整句）+ 关键指标（加粗指标名分点）
    res_sh = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["experiment_result"])
    res_lines = []
    result_summary = str(tech.get("experiment_result_summary") or "").strip()
    if result_summary:
        res_lines.append(f"实验总结：{result_summary}")
    res_lines.append(("BOLD", "关键指标："))
    for point in tech.get("experiment_result_points", []):
        point = str(point or "").strip()
        if point:
            res_lines.append(f"• {point}")
    write_bold_prefix_paragraphs(res_sh.text_frame, res_lines, autofit=False)

    # ── Experiment figures: reflow figure(s) + result text together.
    # 实验图常常比模板给的图片区域更矮更宽（例如柱状图），按比例适配后
    # 图片上下会留白；把图片和结果文字一起重新排布，将这部分留白让给
    # 结果文字，同时把卡片底部本来就空出的一点留白也纳入可用范围，
    # 避免"实验总结 + 关键指标"这类更长的内容顶破卡片下边框。
    exp_figs = _coerce_image_list(
        tech.get("experiment_figures"),
        tech.get("experiment_figure"),
    )
    if not exp_figs:
        raise ValueError(f"{title_text}: missing experiment figure")
    exp_slot = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["experiment_figure"])
    right_card = find_shape_by_id_safe(slide, INSIGHT_SHAPE_IDS["right_card"])
    res_box = _box_from_shape(res_sh)
    if right_card is not None:
        card_bottom = right_card.top + right_card.height - Inches(0.08)
    else:
        card_bottom = res_box[1] + res_box[3]
    extended_result_box = (
        res_box[0],
        res_box[1],
        res_box[2],
        max(1, int(card_bottom) - res_box[1]),
    )
    exp_box = _box_from_shape(exp_slot)
    # 图片区域按结果文字框的左边界和宽度对齐（与卡片留白保持一致），
    # 只保留图片原本的纵向起点；避免图片自身偏窄的占位框把重排后的
    # 文字也一起挤窄。
    aligned_image_box = (res_box[0], exp_box[1], res_box[2], exp_box[3])
    reflow_figure_and_text(
        slide,
        exp_slot,
        exp_figs,
        aligned_image_box,
        res_sh,
        [exp_slot],
        text_box=extended_result_box,
        min_text_fraction=0.55,
    )

    # ── Takeaway
    ins_tbl = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["insight_table"])
    replace_cell_text(
        ins_tbl.table.rows[0].cells[1],
        tech.get("takeaway", ""),
        autofit=False,
    )


def _final_summary_spec(spec):
    return (
        spec.get("final_summary")
        or spec.get("problems_outlook")
        or spec.get("future_trends")
        or {}
    )


def populate_final_summary(slide, spec):
    fs = _final_summary_spec(spec)
    if not fs:
        return

    title = fs.get("title") or "存在问题与未来展望"
    problems = fs.get("problems", [])
    outlook = (
        fs.get("outlook")
        or fs.get("future_outlook")
        or fs.get("trends")
        or []
    )
    # Backward compatibility for older specs that only had future_trends.points.
    if not outlook and fs.get("points"):
        outlook = fs.get("points", [])
    conclusion = fs.get("conclusion") or fs.get("closing") or ""

    problem_lines = [f"问题{i+1}：{text}" for i, text in enumerate(problems)]
    outlook_lines = [f"展望{i+1}：{text}" for i, text in enumerate(outlook)]
    body_lines = problem_lines + outlook_lines

    title_shape = find_shape_by_id_safe(slide, 2)
    body_shape = find_shape_by_id_safe(slide, 4)
    conclusion_shape = find_shape_by_id_safe(slide, 5)

    if title_shape is not None and title_shape.has_text_frame:
        replace_text_preserving_format(
            title_shape.text_frame,
            title,
            line_spacing_pct=SINGLE_SPACING_PCT,
        )
        for para in title_shape.text_frame.paragraphs:
            para.alignment = PP_ALIGN.CENTER
    if body_shape is not None and body_shape.has_text_frame:
        write_bold_prefix_paragraphs(
            body_shape.text_frame,
            body_lines,
            line_spacing_pct=LINE_SPACING_PCT,
        )
    if conclusion_shape is not None and conclusion_shape.has_text_frame:
        replace_text_preserving_format(
            conclusion_shape.text_frame,
            f"判断：{conclusion}" if conclusion else "",
            line_spacing_pct=LINE_SPACING_PCT,
        )


def populate_trend_insight(slide, spec):
    trend = spec.get("trend_insight", {})
    stages = trend.get("stages", [])[:4]
    stage_labels = " → ".join(
        stage.get("label", "") for stage in stages if stage.get("label")
    )
    background_text = "｜".join(
        item
        for item in (
            str(trend.get("time_horizon", "")).strip(),
            stage_labels,
        )
        if item
    )
    _prepare_template_content_slide(
        slide,
        title=trend.get("title") or "技术趋势洞察",
        reference=f"观察窗口： {trend.get('time_horizon', '')}".rstrip(),
        background_label="趋势\n主线",
        background_text=background_text,
        left_heading="技术演进",
        right_heading="代表性进展",
        insight_label="洞察\n观点",
        insight_text=trend.get("viewpoint", ""),
    )

    stage_gap = 0.08
    stage_height = (4.00 - stage_gap * (len(stages) - 1)) / max(1, len(stages))
    stage_colors = ("blue", "cyan", "orange", "red")
    for index, stage in enumerate(stages):
        y = 2.17 + index * (stage_height + stage_gap)
        color = stage_colors[index]
        example = (
            stage.get("example")
            if isinstance(stage.get("example"), dict)
            else {}
        )

        # 同一水平行表示一组因果关系：左侧技术变化，右侧对应的论文或实践。
        _add_textbox(
            slide,
            0.73,
            y,
            5.78,
            stage_height,
            "",
            fill="white" if index % 2 == 0 else "F7F7F7",
            line=None,
        )
        _add_textbox(
            slide,
            0.86,
            y + 0.09,
            0.58,
            0.25,
            stage.get("stage_id", ""),
            size=10,
            color="white",
            bold=True,
            align=PP_ALIGN.CENTER,
            fill=color,
            line=color,
            margin=0.02,
            line_spacing_pct=SINGLE_SPACING_PCT,
        )
        _add_textbox(
            slide,
            1.55,
            y + 0.06,
            2.83,
            0.29,
            stage.get("label", ""),
            size=12.2 if len(stages) <= 3 else 11.1,
            color="blue",
            bold=True,
            margin=0,
            line_spacing_pct=SINGLE_SPACING_PCT,
        )
        _add_textbox(
            slide,
            4.45,
            y + 0.08,
            1.84,
            0.23,
            stage.get("period", ""),
            size=8.7,
            color="muted",
            bold=True,
            align=PP_ALIGN.RIGHT,
            margin=0,
            line_spacing_pct=SINGLE_SPACING_PCT,
        )
        _add_textbox(
            slide,
            0.86,
            y + 0.39,
            5.42,
            max(0.30, stage_height - 0.45),
            stage.get("technical_change") or stage.get("signal", ""),
            size=9.6 if len(stages) <= 3 else 8.6,
            color="text",
            margin=0.02,
            valign=MSO_ANCHOR.TOP,
        )

        _add_textbox(
            slide,
            7.14,
            y,
            5.57,
            stage_height,
            "",
            fill="white" if index % 2 == 0 else "F7F7F7",
            line=None,
        )

        image_height = max(0.50, stage_height - 0.20)
        example_picture = _add_fitted_picture(
            slide,
            example.get("figure", ""),
            7.28,
            y + 0.10,
            1.58,
            image_height,
        )
        if example.get("source_url"):
            example_picture.click_action.hyperlink.address = example.get(
                "source_url"
            )
        example_title_shape = _add_textbox(
            slide,
            9.02,
            y + 0.08,
            3.48,
            0.31,
            _compact_slide_text(example.get("title", ""), 56),
            size=9.7 if len(stages) <= 3 else 8.7,
            color="blue",
            bold=True,
            margin=0,
            valign=MSO_ANCHOR.TOP,
            line_spacing_pct=SINGLE_SPACING_PCT,
        )
        for paragraph in example_title_shape.text_frame.paragraphs:
            for run in paragraph.runs:
                _set_run_style(
                    run,
                    size=8.7 if len(stages) > 3 else 9.7,
                    color="blue",
                    bold=True,
                )
                run.font.underline = False
        meta = "｜".join(
            item
            for item in (
                str(example.get("organization", "")).strip(),
                str(example.get("date", "")).strip(),
            )
            if item
        )
        _add_textbox(
            slide,
            9.02,
            y + 0.40,
            3.48,
            0.14,
            meta,
            size=7.1,
            color="muted",
            margin=0,
            line_spacing_pct=SINGLE_SPACING_PCT,
        )
        _add_textbox(
            slide,
            9.02,
            y + 0.56,
            3.48,
            max(0.25, stage_height - 0.60),
            _compact_slide_text(example.get("fact", ""), 92),
            size=7.9 if len(stages) <= 3 else 7.0,
            color="text",
            margin=0,
            valign=MSO_ANCHOR.TOP,
        )


def _format_table_cell(cell, text, *, size=10, color="text", bold=False, fill=None):
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(fill)
    _format_text_frame(
        cell.text_frame,
        text,
        size=size,
        color=color,
        bold=bold,
        margin=0.05,
        valign=MSO_ANCHOR.MIDDLE,
        line_spacing_pct=SINGLE_SPACING_PCT,
    )


def populate_academic_resources(slide, spec):
    academic = spec.get("academic_resources", {})
    universities = academic.get("universities", [])[:4]
    faculty_projects = academic.get("faculty_projects", [])[:5]
    _prepare_template_content_slide(
        slide,
        title=academic.get("title") or "学术资源洞察",
        reference="资源： 高校实验室主页、教师主页与代表课题",
        background_label="资源\n范围",
        background_text=(
            f"本页汇总 {len(universities)} 所高校及实验室、"
            f"{len(faculty_projects)} 位教师与代表课题，名称可点击查看对应主页。"
        ),
        left_heading="高校资源",
        right_heading="教师与课题",
        insight_label="学术\n观点",
        insight_text=academic.get("viewpoint", ""),
    )

    university_gap = 0.09
    university_height = (
        3.94 - university_gap * (len(universities) - 1)
    ) / max(1, len(universities))
    for index, university in enumerate(universities):
        y = 2.20 + index * (university_height + university_gap)
        body = (
            f"{university.get('name', '')}\n"
            f"{university.get('school_lab', '')}\n"
            f"{university.get('focus', '')}\n"
            f"{university.get('strength', '')}"
        )
        card = _add_textbox(
            slide,
            0.74,
            y,
            5.76,
            university_height,
            body,
            size=9.5 if len(universities) <= 3 else 8.6,
            color="text",
            bold=False,
            fill="white" if index % 2 == 0 else "F7F7F7",
            line=None,
            margin=0.12,
            valign=MSO_ANCHOR.TOP,
        )
        if university.get("url") and card.text_frame.paragraphs:
            for run in card.text_frame.paragraphs[0].runs:
                if run.text:
                    run.hyperlink.address = university.get("url")
        for paragraph_index, paragraph in enumerate(card.text_frame.paragraphs):
            for run in paragraph.runs:
                _set_run_style(
                    run,
                    size=10.1 if paragraph_index == 0 else 8.6,
                    color="blue" if paragraph_index <= 1 else "text",
                    bold=paragraph_index <= 1,
                )

    table_shape = slide.shapes.add_table(
        len(faculty_projects) + 1,
        4,
        Inches(7.15),
        Inches(2.20),
        Inches(5.56),
        Inches(3.94),
    )
    table = table_shape.table
    table.rows[0].height = Inches(0.43)
    faculty_row_height = (3.94 - 0.43) / max(1, len(faculty_projects))
    for row_index in range(1, len(table.rows)):
        table.rows[row_index].height = Inches(faculty_row_height)
    column_widths = (1.00, 1.12, 1.30, 2.14)
    for index, width in enumerate(column_widths):
        table.columns[index].width = Inches(width)
    headers = ("教师", "高校 / 团队", "研究方向", "代表课题")
    for column, header in enumerate(headers):
        _format_table_cell(
            table.cell(0, column),
            header,
            size=8.6,
            color="blue",
            bold=True,
            fill="F0F0F0",
        )
    for row_index, item in enumerate(faculty_projects, start=1):
        faculty_text = "\n".join(
            value
            for value in (
                str(item.get("faculty", "")).strip(),
                str(item.get("title", "")).strip(),
            )
            if value
        )
        project_text = "\n".join(
            value
            for value in (
                str(item.get("project", "")).strip(),
                str(item.get("period", "")).strip(),
            )
            if value
        )
        values = (
            faculty_text,
            item.get("university", ""),
            item.get("research_focus", ""),
            project_text,
        )
        row_fill = "white" if row_index % 2 else "pale_blue"
        for column, value in enumerate(values):
            _format_table_cell(
                table.cell(row_index, column),
                value,
                size=7.5 if column != 3 else 7.2,
                color="text",
                bold=column == 0,
                fill=row_fill,
            )
        for column, url_key in ((0, "profile_url"), (3, "project_url")):
            url = item.get(url_key)
            if url:
                for para in table.cell(row_index, column).text_frame.paragraphs:
                    for run in para.runs:
                        if run.text:
                            run.hyperlink.address = url


def populate_overall_viewpoint(slide, spec):
    viewpoint = spec.get("overall_viewpoint", {})
    _prepare_template_content_slide(
        slide,
        title=viewpoint.get("title") or "技术洞察观点",
        reference="归纳范围： 技术趋势、单项技术与学术资源",
        background_label="核心\n观点",
        background_text=viewpoint.get("viewpoint", ""),
        left_heading="关键判断",
        right_heading="问题与建议",
        insight_label="落地\n建议",
        insight_text=viewpoint.get("suggestion", ""),
    )

    sections = viewpoint.get("sections", [])[:3]
    card_gap = 0.10
    card_height = (3.94 - card_gap * (len(sections) - 1)) / max(1, len(sections))
    card_colors = (
        ("pale_blue", "blue"),
        ("E8FAFA", "cyan"),
        ("pale_orange", "orange"),
    )
    for index, section in enumerate(sections):
        fill_color, line_color = card_colors[index]
        y = 2.20 + index * (card_height + card_gap)
        _add_textbox(
            slide,
            0.74,
            y,
            5.76,
            0.38,
            section.get("title", ""),
            size=11.0,
            color="blue",
            bold=True,
            fill=fill_color,
            line=line_color,
            margin=0.09,
        )
        _add_bullets(
            slide,
            0.74,
            y + 0.38,
            5.76,
            max(0.30, card_height - 0.38),
            section.get("points", [])[:3],
            size=8.5,
            color="text",
            fill="white" if index % 2 == 0 else "F7F7F7",
            line=None,
            margin=0.10,
        )

    unresolved = viewpoint.get("unresolved", [])
    unresolved_text = "\n".join(
        f"• {item}" for item in unresolved[:3]
    )
    _add_textbox(
        slide,
        7.15,
        2.20,
        5.56,
        2.30,
        f"待解问题\n{unresolved_text}",
        size=9.0,
        color="text",
        fill="F7F7F7",
        line=None,
        margin=0.16,
        valign=MSO_ANCHOR.TOP,
    )
    _add_textbox(
        slide,
        7.15,
        4.68,
        5.56,
        1.46,
        (
            "仍需验证\n"
            + "\n".join(f"• {item}" for item in unresolved[3:5])
            if unresolved[3:5]
            else f"建议｜{viewpoint.get('suggestion', '')}"
        ),
        size=9.2,
        color="text",
        bold=True,
        fill="pale_orange",
        line=None,
        margin=0.16,
        valign=MSO_ANCHOR.TOP,
    )


def _resolve_figure_paths(spec, spec_dir):
    resolved = copy.deepcopy(spec)
    for stage in resolved.get("trend_insight", {}).get("stages", []):
        example = stage.get("example") if isinstance(stage.get("example"), dict) else {}
        figure = example.get("figure")
        if figure and not Path(figure).is_absolute():
            example["figure"] = str((spec_dir / figure).resolve())
    for direction in resolved.get("directions", []):
        for tech in direction.get("technologies", []):
            for key in (
                "tech_detail_figures",
                "tech_detail_figure",
                "tech_detail_figure2",
                "experiment_figures",
                "experiment_figure",
            ):
                value = tech.get(key)
                if isinstance(value, list):
                    tech[key] = [
                        str((spec_dir / item).resolve())
                        if item and not Path(item).is_absolute()
                        else item
                        for item in value
                    ]
                elif value and not Path(value).is_absolute():
                    tech[key] = str((spec_dir / value).resolve())
    return resolved


# ── main build ───────────────────────────────────────────────────────────────

def build(spec_path, template_path, out_path):
    _TEMP_IMAGE_PATHS.clear()
    temp_output_path = None
    spec_path = Path(spec_path).resolve()
    out_path = Path(out_path).resolve()
    raw_spec = json.loads(spec_path.read_text(encoding="utf-8"))
    report = validate_spec_data(raw_spec, base_dir=spec_path.parent, phase="build")
    if not report.ok:
        details = "\n".join(f"- {item}" for item in report.errors)
        raise ValueError(f"spec 未通过 build 校验：\n{details}")
    for warning in report.warnings:
        print(f"[WARN] {warning}")

    spec = _resolve_figure_paths(raw_spec, spec_path.parent)
    prs = Presentation(template_path)
    cover_slide = prs.slides[COVER_IDX]
    summary_slide = prs.slides[SUMMARY_IDX]
    techlist_slide = prs.slides[TECHLIST_IDX]
    insight_template = prs.slides[INSIGHT_IDX]
    thanks_slide = prs.slides[THANKS_IDX]

    try:
        tech_pairs = []
        for direction in spec.get("directions", []):
            for tech in direction.get("technologies", []):
                tech_pairs.append((direction, tech))

        # 所有内容页都必须在模板页尚未填充时完成复制。
        # 这样趋势、学术资源、技术洞察观点和每个单项洞察都保留同一套母版与结构形状。
        insight_slides = [insight_template] if tech_pairs else []
        for _ in tech_pairs[1:]:
            insight_slides.append(_new_content_slide(prs))
        trend_slide = _new_content_slide(prs)
        academic_slide = _new_content_slide(prs)
        viewpoint_slide = _new_content_slide(prs)

        populate_cover(cover_slide, spec)
        populate_summary(summary_slide, spec)
        populate_techlist(techlist_slide, spec)
        decorate_summary_slide(summary_slide, spec)
        decorate_techlist_slide(techlist_slide)

        if tech_pairs:
            for slide, (direction, tech) in zip(insight_slides, tech_pairs):
                populate_insight(slide, tech, direction)
        else:
            remove_slide(prs, INSIGHT_IDX)

        populate_trend_insight(trend_slide, spec)
        populate_academic_resources(academic_slide, spec)
        populate_overall_viewpoint(viewpoint_slide, spec)

        ordered_slides = [
            cover_slide,
            trend_slide,
            summary_slide,
            techlist_slide,
            *insight_slides,
            academic_slide,
            viewpoint_slide,
            thanks_slide,
        ]
        for index, slide in enumerate(ordered_slides):
            reorder_slide_to(prs, slide, index)

        normalize_deck_fonts(prs)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            suffix=".pptx",
            prefix=f".{out_path.stem}-",
            dir=out_path.parent,
            delete=False,
        ) as temp_output:
            temp_output_path = Path(temp_output.name)
        prs.save(temp_output_path)
        validate_output_text_fit(temp_output_path)
        fidelity = validate_template_fidelity(
            template_path,
            temp_output_path,
            spec_path,
        )
        if not fidelity.ok:
            details = "\n".join(f"- {item}" for item in fidelity.errors)
            raise ValueError(f"成品未通过模板一致性校验：\n{details}")
        os.replace(temp_output_path, out_path)
        temp_output_path = None
        print(f"[OK] 已生成 {out_path}（{len(prs.slides)} 页）")
    finally:
        if temp_output_path is not None:
            try:
                temp_output_path.unlink()
            except OSError:
                pass
        cleanup_temp_images()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--spec",     required=True)
    ap.add_argument("--template", default=None)
    ap.add_argument("--out",      required=True)
    args = ap.parse_args()
    tmpl = args.template or str(
        Path(__file__).resolve().parent.parent / "assets" / "template.pptx")
    try:
        build(args.spec, tmpl, args.out)
    except Exception as exc:
        print(f"[FAIL] {exc}")
        return 1
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
