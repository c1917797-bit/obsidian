"""校验生成稿是否真正继承指定模板，而不是从空白演示文稿近似重画。"""
from __future__ import annotations

import argparse
import json
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation


INSIGHT_STRUCTURE_IDS = {4, 5, 6, 7, 9, 10, 11, 13}


@dataclass
class FidelityReport:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self):
        return not self.errors


def _master_signature(master):
    return (
        len(master.shapes),
        tuple(layout.name for layout in master.slide_layouts),
    )


def _shape_ids(slide):
    return {shape.shape_id for shape in slide.shapes}


def _technology_count(spec):
    return sum(
        len(direction.get("technologies", []))
        for direction in spec.get("directions", [])
    )


def validate_template_fidelity(template_path, deck_path, spec_path=None):
    template_path = Path(template_path).resolve()
    deck_path = Path(deck_path).resolve()
    report = FidelityReport()

    template = Presentation(template_path)
    deck = Presentation(deck_path)

    if (
        deck.slide_width != template.slide_width
        or deck.slide_height != template.slide_height
    ):
        report.errors.append(
            "页面尺寸与模板不一致；成品很可能由空白演示文稿重新创建。"
        )

    template_master_signatures = Counter(
        _master_signature(master) for master in template.slide_masters
    )
    deck_master_signatures = Counter(
        _master_signature(master) for master in deck.slide_masters
    )
    missing_masters = template_master_signatures - deck_master_signatures
    if missing_masters:
        report.errors.append(
            "模板母版或版式未完整继承；禁止用 Presentation.create 或空白页重画。"
        )

    allowed_layouts = {
        slide.slide_layout.name
        for slide in template.slides
    }
    unknown_layouts = sorted(
        {
            slide.slide_layout.name
            for slide in deck.slides
            if slide.slide_layout.name not in allowed_layouts
        }
    )
    if unknown_layouts:
        report.errors.append(
            "成品使用了模板之外的版式："
            + "、".join(unknown_layouts)
        )

    if spec_path:
        spec = json.loads(
            Path(spec_path).resolve().read_text(encoding="utf-8")
        )
        tech_count = _technology_count(spec)
        expected_count = tech_count + 7
        if len(deck.slides) != expected_count:
            report.errors.append(
                f"页数与固定叙事顺序不一致：应为 {expected_count} 页，"
                f"实际为 {len(deck.slides)} 页。"
            )
        elif len(deck.slides) >= 7:
            insight_like_indices = (
                [1]
                + list(range(4, 4 + tech_count))
                + [4 + tech_count, 5 + tech_count]
            )
            for index in insight_like_indices:
                missing_ids = (
                    INSIGHT_STRUCTURE_IDS - _shape_ids(deck.slides[index])
                )
                if missing_ids:
                    report.errors.append(
                        f"第 {index + 1} 页缺少模板结构形状 "
                        f"{sorted(missing_ids)}；不得删除标题、双栏框、"
                        "背景带或观点带后另行重画。"
                    )

    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--deck", required=True)
    parser.add_argument("--spec")
    args = parser.parse_args()

    report = validate_template_fidelity(
        args.template,
        args.deck,
        args.spec,
    )
    for warning in report.warnings:
        print(f"[WARN] {warning}")
    if not report.ok:
        for error in report.errors:
            print(f"[FAIL] {error}")
        return 1
    print("[OK] 成品继承了模板页面尺寸、母版、版式与关键结构。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
