"""Build a technology-insight deck from a JSON spec, using the fixed template."""
from __future__ import annotations

import argparse
import copy
import json
import os
import re
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt
from pptx.oxml.ns import qn
from lxml import etree

try:
    from PIL import Image
except ImportError:  # pragma: no cover - surfaced at runtime with a clear error.
    Image = None

COVER_IDX   = 0
SUMMARY_IDX = 1
TECHLIST_IDX = 2
INSIGHT_IDX  = 3
THANKS_IDX   = 4

INSIGHT_SHAPE_IDS = {
    "title":             5,
    "reference":        13,
    "background_table": 10,
    "tech_detail_figure":  15,
    "tech_detail_figure2": 21,
    "tech_detail_text":    23,
    "experiment_figure":   17,
    "experiment_method":   25,
    "experiment_result":   26,
    "insight_table":       11,
}

LINE_SPACING_PCT  = 120000   # 1.2×
SINGLE_SPACING_PCT = 100000  # 1.0×
_a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
_r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_TEMP_IMAGE_PATHS = []


# ── shape lookup ────────────────────────────────────────────────────────────

def find_shape_by_id(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError(f"shape_id {shape_id} not found")

def find_shape_by_id_safe(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    return None

def remove_shape(slide, shape_id):
    sh = find_shape_by_id_safe(slide, shape_id)
    if sh is not None:
        sh._element.getparent().remove(sh._element)


# ── line-spacing & autofit ──────────────────────────────────────────────────

def _set_line_spacing(text_frame, pct):
    for para in text_frame.paragraphs:
        pPr = para._p.find(qn("a:pPr"))
        if pPr is None:
            pPr = etree.SubElement(para._p, qn("a:pPr"))
            para._p.insert(0, pPr)
        lnSpc = pPr.find(qn("a:lnSpc"))
        if lnSpc is None:
            lnSpc = etree.SubElement(pPr, qn("a:lnSpc"))
        for ch in list(lnSpc):
            lnSpc.remove(ch)
        sp = etree.SubElement(lnSpc, qn("a:spcPct"))
        sp.set("val", str(pct))

def apply_line_spacing(text_frame, pct=LINE_SPACING_PCT):
    _set_line_spacing(text_frame, pct)

def apply_single_spacing(text_frame):
    _set_line_spacing(text_frame, SINGLE_SPACING_PCT)

def apply_line_spacing_cell(cell, pct=LINE_SPACING_PCT):
    _set_line_spacing(cell.text_frame, pct)

def enable_text_autofit(text_frame):
    """Set normAutofit so PowerPoint shrinks the font to fit rather than clip."""
    txBody = text_frame._txBody
    bodyPr = txBody.find(qn("a:bodyPr"))
    if bodyPr is None:
        return
    for tag in [qn("a:noAutofit"), qn("a:normAutofit"), qn("a:spAutoFit")]:
        for el in bodyPr.findall(tag):
            bodyPr.remove(el)
    etree.SubElement(bodyPr, qn("a:normAutofit"))


def _clean_text_lines(value):
    lines = [line.strip() for line in str(value or "").splitlines()]
    return [line for line in lines if line]


# ── text writing helpers ────────────────────────────────────────────────────

def _pick_style_run(text_frame):
    candidates = []
    for para in text_frame.paragraphs:
        for run in para.runs:
            candidates.append(run)
    if not candidates:
        return None
    def is_bold(r):
        try: return bool(r.font.bold)
        except: return False
    pool = [r for r in candidates if not is_bold(r)] or candidates
    pool.sort(key=lambda r: len(r.text or ""), reverse=True)
    return pool[0]._r


def replace_text_preserving_format(text_frame, new_text,
                                   line_spacing_pct=LINE_SPACING_PCT,
                                   autofit=True):
    if text_frame is None:
        return
    style_r   = _pick_style_run(text_frame)
    style_xml = copy.deepcopy(style_r) if style_r is not None else None
    body = text_frame._txBody
    for p in body.findall(f"{{{_a_ns}}}p"):
        body.remove(p)
    for line in (_clean_text_lines(new_text) or [""]):
        p_el = body.makeelement(qn("a:p"), {})
        body.append(p_el)
        if style_xml is not None:
            r_el = copy.deepcopy(style_xml)
            t_el = r_el.find(qn("a:t"))
            if t_el is not None:
                t_el.text = line
            p_el.append(r_el)
        else:
            r_el = p_el.makeelement(qn("a:r"), {})
            t_el = r_el.makeelement(qn("a:t"), {})
            t_el.text = line
            r_el.append(t_el)
            p_el.append(r_el)
    _set_line_spacing(text_frame, line_spacing_pct)
    if autofit:
        enable_text_autofit(text_frame)


def replace_cell_text(cell, new_text, line_spacing_pct=LINE_SPACING_PCT):
    replace_text_preserving_format(cell.text_frame, new_text,
                                   line_spacing_pct=line_spacing_pct)


def write_bold_prefix_paragraphs(text_frame, lines,
                                  line_spacing_pct=LINE_SPACING_PCT):
    """Write paragraphs where text before fullwidth '：' is bold.
    lines can be plain strings or ('BOLD'/'NORMAL', text) tuples.
    """
    style_r   = _pick_style_run(text_frame)
    style_xml = copy.deepcopy(style_r) if style_r is not None else None
    body = text_frame._txBody
    for p in body.findall(f"{{{_a_ns}}}p"):
        body.remove(p)

    def make_run(tmpl, content, bold):
        if tmpl is not None:
            r_el = copy.deepcopy(tmpl)
            rPr  = r_el.find(qn("a:rPr"))
            if rPr is None:
                rPr = etree.SubElement(r_el, qn("a:rPr"))
                r_el.insert(0, rPr)
            if bold:
                rPr.set("b", "1")
            else:
                rPr.attrib.pop("b", None)
            t_el = r_el.find(qn("a:t"))
            if t_el is None:
                t_el = etree.SubElement(r_el, qn("a:t"))
            t_el.text = content
        else:
            r_el = etree.Element(qn("a:r"))
            rPr  = etree.SubElement(r_el, qn("a:rPr"))
            if bold:
                rPr.set("b", "1")
            t_el = etree.SubElement(r_el, qn("a:t"))
            t_el.text = content
        return r_el

    cleaned_lines = []
    for line in lines:
        if isinstance(line, tuple):
            mode, content = line
            content = str(content or "").strip()
            if content:
                cleaned_lines.append((mode, content))
        else:
            content = str(line or "").strip()
            if content:
                cleaned_lines.append(content)

    for line in cleaned_lines:
        p_el = body.makeelement(qn("a:p"), {})
        body.append(p_el)
        if isinstance(line, tuple):
            mode, content = line
            p_el.append(make_run(style_xml, content, mode == "BOLD"))
        else:
            sep = "\uff1a" if "\uff1a" in line else (
                  ":" if ":" in line and line.index(":") < 28 else None)
            if sep:
                idx    = line.index(sep) + len(sep)
                prefix = line[:idx]
                rest   = line[idx:]
                if prefix: p_el.append(make_run(style_xml, prefix, True))
                if rest:   p_el.append(make_run(style_xml, rest,   False))
            else:
                p_el.append(make_run(style_xml, line, False))

    _set_line_spacing(text_frame, line_spacing_pct)
    enable_text_autofit(text_frame)


# ── picture replacement (z-order safe, no shared-blob issue) ─────────────────

def _box_from_shape(shape):
    return shape.left, shape.top, shape.width, shape.height


def _union_box(shapes):
    left = min(sh.left for sh in shapes)
    top = min(sh.top for sh in shapes)
    right = max(sh.left + sh.width for sh in shapes)
    bottom = max(sh.top + sh.height for sh in shapes)
    return left, top, right - left, bottom - top


def _image_ratio(path):
    if Image is None:
        raise RuntimeError("Pillow is required for aspect-ratio-safe image placement.")
    with Image.open(path) as img:
        return img.width / img.height


def _prepare_image_for_box(image_path, width, height):
    """Center-crop image to the target box ratio so PowerPoint does not distort it."""
    if Image is None:
        raise RuntimeError("Pillow is required for aspect-ratio-safe image placement.")
    target_ratio = float(width) / float(height)
    img = Image.open(image_path).convert("RGB")
    src_ratio = img.width / img.height
    if abs(src_ratio - target_ratio) < 0.01:
        prepared = img
    elif src_ratio > target_ratio:
        new_w = int(img.height * target_ratio)
        left = max(0, (img.width - new_w) // 2)
        prepared = img.crop((left, 0, left + new_w, img.height))
    else:
        new_h = int(img.width / target_ratio)
        top = max(0, (img.height - new_h) // 2)
        prepared = img.crop((0, top, img.width, top + new_h))

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
    tmp.close()
    prepared.save(tmp.name)
    _TEMP_IMAGE_PATHS.append(tmp.name)
    return tmp.name


def _remove_shape_elements(shapes):
    for sh in shapes:
        if sh is not None and sh._element.getparent() is not None:
            sh._element.getparent().remove(sh._element)


def _add_picture_at(slide, anchor_shape, image_path, left, top, width, height,
                    remove_shapes):
    if not os.path.exists(image_path):
        raise FileNotFoundError(f"Image not found: {image_path}")
    spTree = anchor_shape._element.getparent()
    z_idx = list(spTree).index(anchor_shape._element)
    prepared_path = _prepare_image_for_box(image_path, width, height)
    _remove_shape_elements(remove_shapes)
    new_sh = slide.shapes.add_picture(
        prepared_path, left, top, width=width, height=height)
    new_el = new_sh._element
    spTree.remove(new_el)
    spTree.insert(z_idx, new_el)
    return new_sh


def replace_picture(slide, shape_id, new_image_path):
    """Replace picture while preserving z-order and image aspect ratio.

    Uses remove-then-add, but re-inserts the new shape element at the
    same spTree index as the original. The source image is center-cropped
    to the placeholder ratio before insertion, so PowerPoint never stretches it.
    Also creates a fresh image Part per call, avoiding the shared-blob
    issue that arises when slides are cloned and all reference the same Part.
    """
    pic = find_shape_by_id(slide, shape_id)
    left, top, width, height = _box_from_shape(pic)
    return _add_picture_at(slide, pic, new_image_path, left, top, width, height,
                           [pic])


def _coerce_image_list(value, *fallbacks):
    paths = []
    if isinstance(value, list):
        paths.extend(value)
    elif value:
        paths.append(value)
    paths.extend(f for f in fallbacks if f)
    seen = set()
    result = []
    for path in paths:
        if path and path not in seen:
            seen.add(path)
            result.append(path)
    return result


def _split_box_for_images(image_paths, left, top, width, height):
    if len(image_paths) <= 1:
        return [(left, top, width, height)]
    gap = int(min(width, height) * 0.035)
    ratios = [_image_ratio(path) for path in image_paths[:2]]
    side_ratio = ((width - gap) / 2) / height
    stack_ratio = width / ((height - gap) / 2)

    def penalty(target):
        return sum(abs(r - target) / max(r, target) for r in ratios)

    if penalty(stack_ratio) < penalty(side_ratio):
        half_h = int((height - gap) / 2)
        return [
            (left, top, width, half_h),
            (left, top + half_h + gap, width, height - half_h - gap),
        ]
    half_w = int((width - gap) / 2)
    return [
        (left, top, half_w, height),
        (left + half_w + gap, top, width - half_w - gap, height),
    ]


def replace_picture_group(slide, anchor_shape, image_paths, box, remove_shapes):
    image_paths = image_paths[:2]
    left, top, width, height = box
    boxes = _split_box_for_images(image_paths, left, top, width, height)
    spTree = anchor_shape._element.getparent()
    z_idx = list(spTree).index(anchor_shape._element)
    _remove_shape_elements(remove_shapes)
    inserted = []
    for path, (x, y, w, h) in zip(image_paths, boxes):
        if not os.path.exists(path):
            raise FileNotFoundError(f"Image not found: {path}")
        prepared_path = _prepare_image_for_box(path, w, h)
        new_sh = slide.shapes.add_picture(prepared_path, x, y, width=w, height=h)
        inserted.append(new_sh._element)
    for el in inserted:
        spTree.remove(el)
    for offset, el in enumerate(inserted):
        spTree.insert(z_idx + offset, el)


def cleanup_temp_images():
    for path in list(_TEMP_IMAGE_PATHS):
        try:
            os.remove(path)
        except OSError:
            pass
    _TEMP_IMAGE_PATHS.clear()


# ── title auto-shrink ────────────────────────────────────────────────────────

def auto_fit_title(title_shape, title_text):
    """Shrink font so the title stays on one line (threshold: 65 vis-width units)."""
    visible = sum(2 if ord(c) > 127 else 1 for c in title_text)
    if visible <= 65:
        return
    new_pt = max(13, int(22 * 65 / visible))
    for para in title_shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(new_pt)


def infer_reference_url(reference, explicit_url=None):
    if explicit_url:
        return explicit_url
    ref = reference or ""
    url_match = re.search(r"https?://[^\s,;]+", ref)
    if url_match:
        return url_match.group(0)
    arxiv_match = re.search(r"arxiv\s*:\s*(\d{4}\.\d{4,5})(?:v\d+)?", ref, re.I)
    if arxiv_match:
        return f"https://arxiv.org/abs/{arxiv_match.group(1)}"
    doi_match = re.search(r"doi\s*:\s*([^\s,;]+)", ref, re.I)
    if doi_match:
        return f"https://doi.org/{doi_match.group(1)}"
    return None


def apply_text_hyperlink(text_frame, url):
    if not url:
        return
    for para in text_frame.paragraphs:
        for run in para.runs:
            if run.text:
                run.hyperlink.address = url


# ── slide management ─────────────────────────────────────────────────────────

def duplicate_slide(prs, source_idx):
    source   = prs.slides[source_idx]
    new_sl   = prs.slides.add_slide(source.slide_layout)
    src_tree = source.shapes._spTree
    dst_tree = new_sl.shapes._spTree
    for shp in list(dst_tree):
        if shp.tag.split("}", 1)[-1] in ("sp","pic","graphicFrame","grpSp","cxnSp"):
            dst_tree.remove(shp)
    for shp in src_tree:
        if shp.tag.split("}", 1)[-1] in ("sp","pic","graphicFrame","grpSp","cxnSp"):
            dst_tree.append(copy.deepcopy(shp))
    src_part = source.part
    dst_part = new_sl.part
    for blip in dst_tree.iter(f"{{{_a_ns}}}blip"):
        embed = blip.get(f"{{{_r_ns}}}embed")
        if embed and embed in src_part.rels:
            rel     = src_part.rels[embed]
            new_rid = dst_part.relate_to(rel.target_part, rel.reltype)
            blip.set(f"{{{_r_ns}}}embed", new_rid)
    for hlink in dst_tree.iter(f"{{{_a_ns}}}hlinkClick"):
        rid = hlink.get(f"{{{_r_ns}}}id")
        if rid and rid in src_part.rels:
            rel = src_part.rels[rid]
            new_rid = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True) \
                      if rel.is_external else dst_part.relate_to(rel.target_part, rel.reltype)
            hlink.set(f"{{{_r_ns}}}id", new_rid)
    return new_sl


def reorder_slide_to(prs, slide, new_index):
    sldIdLst = prs.slides._sldIdLst
    target   = slide.part.partname
    for sldId in sldIdLst:
        rId = sldId.get(f"{{{_r_ns}}}id")
        if prs.part.related_part(rId).partname == target:
            sldIdLst.remove(sldId)
            sldIdLst.insert(new_index, sldId)
            return
    raise RuntimeError("Could not locate sldId for slide")


def remove_slide(prs, slide_idx):
    sldIdLst = prs.slides._sldIdLst
    sldId    = list(sldIdLst)[slide_idx]
    rId      = sldId.get(f"{{{_r_ns}}}id")
    prs.part.drop_rel(rId)
    sldIdLst.remove(sldId)


# ── per-slide population ─────────────────────────────────────────────────────

def populate_cover(slide, spec):
    title  = spec.get("title", "").replace("\n", " ")   # collapse to one line
    dept   = spec.get("department", "")
    author = spec.get("author", "")
    date   = spec.get("date", "")
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if "技术洞察" in txt and ("xxxx" in txt or txt.strip() == "技术洞察"):
            # Cover title: single spacing, bold
            replace_text_preserving_format(sh.text_frame, title,
                                           line_spacing_pct=SINGLE_SPACING_PCT,
                                           autofit=False)
            # Make every run bold
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    run.font.bold = True
        elif "部门" in txt and "作者" in txt:
            replace_text_preserving_format(
                sh.text_frame,
                f"部门：{dept}\n作者：{author}\n日期：{date}",
                line_spacing_pct=SINGLE_SPACING_PCT, autofit=False)


def populate_summary(slide, spec):
    topic      = spec.get("topic", "")
    directions = spec.get("directions", [])
    overall    = spec.get("overall_summary", "")
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if "本次洞察主要围绕" in txt:
            replace_text_preserving_format(
                sh.text_frame,
                f"本次洞察主要围绕 {topic}，针对如下几个方向展开洞察：")
        elif "方向1：" in txt or "方向 1：" in txt:
            lines = [f"方向{i+1}：{d['summary']}"
                     for i, d in enumerate(directions)]
            write_bold_prefix_paragraphs(sh.text_frame, lines)
        elif "总结：" in txt:
            write_bold_prefix_paragraphs(sh.text_frame, [f"总结：{overall}"])


def populate_techlist(slide, spec):
    tbl_sh = next((sh for sh in slide.shapes if sh.has_table), None)
    if not tbl_sh:
        return
    plan = []
    for d in spec.get("directions", []):
        cat   = d.get("category_label") or d.get("name", "")
        techs = d.get("technologies", [])
        for j, t in enumerate(techs):
            plan.append((cat if j == 0 else "", t.get("name",""), t.get("brief","")))
    tbl_el  = tbl_sh.table._tbl
    hr      = 1   # header rows
    cur     = len(tbl_sh.table.rows) - hr
    target  = len(plan) if plan else 6
    if target > cur:
        last = tbl_el.findall(qn("a:tr"))[-1]
        for _ in range(target - cur):
            tbl_el.append(copy.deepcopy(last))
    elif target < cur:
        for r in tbl_el.findall(qn("a:tr"))[hr + target:]:
            tbl_el.remove(r)
    table = tbl_sh.table
    for i, (cat, name, brief) in enumerate(plan):
        row = table.rows[hr + i]
        replace_cell_text(row.cells[0], cat)
        replace_cell_text(row.cells[1], name)
        replace_cell_text(row.cells[2], brief)
    for i in range(len(plan), len(table.rows) - hr):
        for c in range(3):
            replace_cell_text(table.rows[hr + i].cells[c], "")


def populate_insight(slide, tech):
    # ── Title (auto-shrink to one line, single spacing, bold preserved)
    title_text  = tech.get("title", "")
    title_shape = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["title"])
    replace_text_preserving_format(title_shape.text_frame, title_text,
                                   line_spacing_pct=SINGLE_SPACING_PCT,
                                   autofit=False)
    auto_fit_title(title_shape, title_text)

    # ── Reference (just below title, single spacing)
    ref_shape = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["reference"])
    ref       = tech.get("reference", "")
    replace_text_preserving_format(ref_shape.text_frame,
                                   f"参考： {ref}" if ref else "参考：",
                                   line_spacing_pct=SINGLE_SPACING_PCT,
                                   autofit=False)
    apply_text_hyperlink(
        ref_shape.text_frame,
        infer_reference_url(ref, tech.get("reference_url")),
    )

    # ── Background table cell
    bg_tbl = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["background_table"])
    replace_cell_text(bg_tbl.table.rows[0].cells[1], tech.get("background", ""))

    # ── Tech-detail bullets (bold key terms)
    det_sh = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["tech_detail_text"])
    write_bold_prefix_paragraphs(det_sh.text_frame,
                                 tech.get("tech_detail_points", []))

    # ── Tech-detail figures. One image spans both slots; two images fill both.
    detail_figs = _coerce_image_list(
        tech.get("tech_detail_figures"),
        tech.get("tech_detail_figure"),
        tech.get("tech_detail_figure2"),
    )
    if not detail_figs:
        raise ValueError(f"{title_text}: missing tech detail figure")
    fig_slot_1 = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["tech_detail_figure"])
    fig_slot_2 = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["tech_detail_figure2"])
    if len(detail_figs) == 1:
        replace_picture_group(
            slide, fig_slot_1, detail_figs, _union_box([fig_slot_1, fig_slot_2]),
            [fig_slot_1, fig_slot_2],
        )
    else:
        replace_picture(slide, INSIGHT_SHAPE_IDS["tech_detail_figure"], detail_figs[0])
        replace_picture(slide, INSIGHT_SHAPE_IDS["tech_detail_figure2"], detail_figs[1])

    # ── Experiment method (bold header)
    meth_sh    = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["experiment_method"])
    meth_lines = []
    if tech.get("experiment_method_header"):
        meth_lines.append(("BOLD", tech["experiment_method_header"]))
    meth_lines.extend(tech.get("experiment_method_points", []))
    write_bold_prefix_paragraphs(meth_sh.text_frame, meth_lines)

    # ── Experiment figures. Supports one wide figure or two arranged figures.
    exp_figs = _coerce_image_list(
        tech.get("experiment_figures"),
        tech.get("experiment_figure"),
    )
    if not exp_figs:
        raise ValueError(f"{title_text}: missing experiment figure")
    exp_slot = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["experiment_figure"])
    replace_picture_group(slide, exp_slot, exp_figs, _box_from_shape(exp_slot), [exp_slot])

    # ── Experiment results
    res_sh    = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["experiment_result"])
    res_lines = []
    if tech.get("experiment_result_header"):
        res_lines.append(("BOLD", tech["experiment_result_header"]))
    res_lines.extend(tech.get("experiment_result_points", []))
    write_bold_prefix_paragraphs(res_sh.text_frame, res_lines)

    # ── Takeaway
    ins_tbl = find_shape_by_id(slide, INSIGHT_SHAPE_IDS["insight_table"])
    replace_cell_text(ins_tbl.table.rows[0].cells[1], tech.get("takeaway", ""))


def _final_summary_spec(spec):
    return (
        spec.get("final_summary")
        or spec.get("problems_outlook")
        or spec.get("future_trends")
        or {}
    )


def populate_final_summary(slide, spec):
    fs = _final_summary_spec(spec)
    if not fs:
        return

    title = fs.get("title") or "存在问题与未来展望"
    problems = fs.get("problems", [])
    outlook = (
        fs.get("outlook")
        or fs.get("future_outlook")
        or fs.get("trends")
        or []
    )
    # Backward compatibility for older specs that only had future_trends.points.
    if not outlook and fs.get("points"):
        outlook = fs.get("points", [])
    conclusion = fs.get("conclusion") or fs.get("closing") or ""

    problem_lines = [f"问题{i+1}：{text}" for i, text in enumerate(problems)]
    outlook_lines = [f"展望{i+1}：{text}" for i, text in enumerate(outlook)]
    body_lines = problem_lines + outlook_lines

    title_shape = find_shape_by_id_safe(slide, 2)
    body_shape = find_shape_by_id_safe(slide, 4)
    conclusion_shape = find_shape_by_id_safe(slide, 5)

    if title_shape is not None and title_shape.has_text_frame:
        replace_text_preserving_format(
            title_shape.text_frame,
            title,
            line_spacing_pct=SINGLE_SPACING_PCT,
        )
        for para in title_shape.text_frame.paragraphs:
            para.alignment = PP_ALIGN.CENTER
    if body_shape is not None and body_shape.has_text_frame:
        write_bold_prefix_paragraphs(
            body_shape.text_frame,
            body_lines,
            line_spacing_pct=LINE_SPACING_PCT,
        )
    if conclusion_shape is not None and conclusion_shape.has_text_frame:
        replace_text_preserving_format(
            conclusion_shape.text_frame,
            f"判断：{conclusion}" if conclusion else "",
            line_spacing_pct=LINE_SPACING_PCT,
        )


# ── main build ───────────────────────────────────────────────────────────────

def build(spec_path, template_path, out_path):
    _TEMP_IMAGE_PATHS.clear()
    spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
    prs  = Presentation(template_path)

    populate_cover(prs.slides[COVER_IDX], spec)
    populate_summary(prs.slides[SUMMARY_IDX], spec)
    populate_techlist(prs.slides[TECHLIST_IDX], spec)

    techs = []
    for d in spec.get("directions", []):
        techs.extend(d.get("technologies", []))

    if not techs:
        remove_slide(prs, INSIGHT_IDX)
    else:
        # Clone all insight slides BEFORE populating, so every clone starts
        # from the pristine template images.
        cloned = [prs.slides[INSIGHT_IDX]]
        for _ in techs[1:]:
            cloned.append(duplicate_slide(prs, INSIGHT_IDX))
        for sl, tech in zip(cloned, techs):
            populate_insight(sl, tech)

        # Optional final summary slide (problems + future outlook), before thanks.
        if _final_summary_spec(spec):
            fs_sl = duplicate_slide(prs, SUMMARY_IDX)
            populate_final_summary(fs_sl, spec)

        # Move thanks slide to end
        thanks = None
        for s in prs.slides:
            if s.slide_layout.name == "End page":
                thanks = s; break
            for sh in s.shapes:
                if sh.has_text_frame and "Thank you" in sh.text_frame.text:
                    thanks = s; break
            if thanks: break
        if thanks:
            reorder_slide_to(prs, thanks, len(prs.slides) - 1)

    prs.save(out_path)
    cleanup_temp_images()
    print(f"Saved {out_path}  ({len(prs.slides)} slides)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--spec",     required=True)
    ap.add_argument("--template", default=None)
    ap.add_argument("--out",      required=True)
    args = ap.parse_args()
    tmpl = args.template or str(
        Path(__file__).resolve().parent.parent / "assets" / "template.pptx")
    build(args.spec, tmpl, args.out)

if __name__ == "__main__":
    main()
